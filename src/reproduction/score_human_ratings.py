#!/usr/bin/env python3
"""Validate a returned reader workbook and summarize each annotation question."""
from __future__ import annotations

import argparse
import csv
from collections import Counter, defaultdict
import json
from pathlib import Path
import re

from openpyxl import load_workbook

from src.reproduction.config import SOURCE_PAPER_DIR as EXPERIMENT_DIR

PACKAGE = EXPERIMENT_DIR / 'validation'
QUESTIONS = ("relation_asserted", "trigger_correct", "carrier_correct", "endpoints_correct")
ANSWERS = {"Yes": "yes", "No": "no", "Uncertain": "uncertain", None: "unanswered", "": "unanswered"}


def score(path: Path, package: Path = PACKAGE) -> dict:
    with (package / "validation-items.csv").open(encoding="utf-8", newline="") as stream:
        source_rows = list(csv.DictReader(stream))
    source = {row["item_id"]: row for row in source_rows}
    if len(source) != len(source_rows):
        raise ValueError("Duplicate item IDs in validation-items.csv")
    ws = load_workbook(path, data_only=False)["Rating Form"]
    ratings = {}
    groups = defaultdict(lambda: {q: Counter() for q in QUESTIONS})
    for cells in ws.iter_rows(min_row=2, values_only=True):
        iid = cells[0]
        if iid in (None, "Example 1", "Example 2"):
            continue
        if iid not in source or iid in ratings:
            raise ValueError(f"Unknown or duplicate item ID: {iid}")
        original = source[iid]["sentence"].replace("<<", "").replace(">>", "")
        if cells[1] != original:
            raise ValueError(f"Source sentence changed: {iid}")
        pair = source[iid]["marked_pair"].replace(", ", "\n")
        if cells[2] != pair:
            raise ValueError(f"Marked pair changed: {iid}")
        trigger = source[iid]["reader_says_trigger"]
        trigger = "None" if trigger in ("", "(none)") else trigger
        carriers = list(dict.fromkeys(
            word.strip() for word in re.findall(r"([^,]+?)\s*\([^)]+\)", source[iid]["reader_says_carriers"])
        ))
        if cells[3] != trigger or cells[4] != ("\n".join(carriers) or "None"):
            raise ValueError(f"Reader proposal changed: {iid}")
        answers = {}
        for q, value in zip(QUESTIONS, cells[5:9]):
            if value not in ANSWERS:
                raise ValueError(f"Invalid answer at {iid}/{q}: {value!r}")
            answers[q] = ANSWERS[value]
        ratings[iid] = answers  # Returned free-text notes are private and never exported.
        for group in ("all", f"corpus:{source[iid]['corpus']}", f"carrier:{source[iid]['carrier_class']}"):
            for q, answer in answers.items():
                groups[group][q][answer] += 1
    if set(ratings) != set(source):
        raise ValueError(f"Missing item IDs: {sorted(set(source) - set(ratings))}")
    summary = {}
    for group, questions in sorted(groups.items()):
        summary[group] = {}
        for q, counts in questions.items():
            definite = counts["yes"] + counts["no"]
            summary[group][q] = {
                **{answer: counts[answer] for answer in ("yes", "no", "uncertain", "unanswered")},
                "yes_fraction_definite": counts["yes"] / definite if definite else None,
            }
    return {"n_items": len(ratings), "summary": summary, "ratings": ratings}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--package", type=Path, default=PACKAGE)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    result = score(args.workbook, args.package)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Validated {result['n_items']} items; wrote {args.out}")


if __name__ == "__main__":
    main()
