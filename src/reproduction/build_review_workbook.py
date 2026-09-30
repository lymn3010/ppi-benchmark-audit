#!/usr/bin/env python3
"""Build the supplementary workbook from saved pattern and carrier observations."""
from __future__ import annotations

from src.reproduction.config import PAPER_DIR as EXPERIMENT_DIR

import argparse
import csv
from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parents[2]

CPR = EXPERIMENT_DIR / 'results/cross-corpus/pattern-rates.csv'
CROSS = EXPERIMENT_DIR / 'results/cross-corpus/fisher-tests.csv'
WITHIN = EXPERIMENT_DIR / 'results/within-corpus-summary.csv'
CORE = EXPERIMENT_DIR / 'results/within-corpus-physical-cells.csv'

ASSESSABLE_SUPPORT = 5  # cells with fewer firings are the excluded long tail

HEADER_FILL = PatternFill("solid", fgColor="FF1F4E5F")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFFFF", size=11)
SIG_FILL = PatternFill("solid", fgColor="FFFDE9D9")       # significant cross-corpus row
SUSPECT_FILL = PatternFill("solid", fgColor="FFFDE0E0")   # carrier token that looks like parse noise
TITLE_FONT = Font(name="Calibri", bold=True, size=13)
WRAP = Alignment(wrap_text=True, vertical="top")


def _read(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def _num(value: str) -> float | int | str:
    """Coerce a ledger string to int/float where possible, else pass through."""
    if value is None or value == "":
        return ""
    try:
        f = float(value)
        return int(f) if f.is_integer() and "." not in value and "e" not in value.lower() else f
    except ValueError:
        return value


def _split_pattern(key: str) -> list[str]:
    """pattern key: family|shape|trigger|case_marker|voice (case_marker '_' = none)."""
    parts = key.split("|")
    parts += [""] * (5 - len(parts))
    fam, shape, trig, case, voice = parts[:5]
    return [fam, shape, trig, "" if case == "_" else case, voice]


_ITALIC = InlineFont(i=True)


def _readable_pattern(fam: str, shape: str, trig: str, case: str, voice: str) -> CellRichText:
    """Render a pattern key as an A/B phrase; inferred words are italic."""
    trig = (trig or "").strip()
    case = (case or "").strip()

    def it(text: str) -> TextBlock:
        return TextBlock(_ITALIC, text)

    if fam == "relational":
        # "A is a <role> of B" -- "is a"/"of" inferred (italic); role noun literal.
        if trig in ("", "is-a", "other"):
            return CellRichText("A ", it("is"), " B")
        article = "an" if trig[:1].lower() in "aeiou" else "a"
        link = case if case else "of"
        return CellRichText("A ", it(f"is {article} "), trig, it(f" {link} "), "B")

    if shape == "AB":
        if voice == "nominalized":
            # event nominal: "<trigger> of A and B" (of/and inferred)
            return CellRichText(trig or "interaction", it(" of "), "A", it(" and "), "B")
        if case:  # "A interact with B" -- preposition is a real token
            return CellRichText("A ", trig, f" {case} ", "B")
        return CellRichText("A ", trig, " B")

    if shape == "AA":
        # both proteins on the same side of the trigger (coordination/complex)
        if voice == "nominalized":
            return CellRichText(trig or "interaction", it(" of "), "A", it(" and "), "B")
        return CellRichText("A", it(" and "), "B ", trig)

    if shape == "NESTED":
        # control construction: A and B are arguments of a nested inner event
        tail = f" {case} " if case else " "
        return CellRichText("A ", trig, tail, "B ", it("(nested control)"))

    return CellRichText("A ", trig, " B")


def _carrier_suspect(word: str) -> bool:
    """A carrier head that is unlikely to be a real noun -> eyeball candidate."""
    w = (word or "").strip()
    if len(w) <= 1:
        return True
    return not any(ch.isalpha() for ch in w)


def _style_header(ws, ncols: int, title: str) -> None:
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=2, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A3"


def _write_grid(ws, headers, rows, title, *, widths=None, row_fill=None) -> None:
    _style_header(ws, len(headers), title)
    for c, h in enumerate(headers, start=1):
        ws.cell(row=2, column=c, value=h)
    for r, record in enumerate(rows, start=3):
        for c, value in enumerate(record, start=1):
            ws.cell(row=r, column=c, value=value)
        if row_fill:
            fill = row_fill(record)
            if fill:
                for c in range(1, len(headers) + 1):
                    ws.cell(row=r, column=c).fill = fill
    ncols = len(headers)
    ws.auto_filter.ref = f"A2:{get_column_letter(ncols)}{max(2, len(rows) + 2)}"
    widths = widths or [16] * ncols
    for c, w in enumerate(widths[:ncols], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w


def build(out_path: Path) -> Path:
    cpr = _read(CPR)
    by_dim: dict[str, list[dict[str, str]]] = {}
    for row in cpr:
        by_dim.setdefault(row["dimension"], []).append(row)

    wb = Workbook()

    # ---- README ----
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 110
    lines = [
        ("Pattern / carrier review workbook -- APBC PPI corpus audit", TITLE_FONT),
        ("", None),
        ("What this is: every pattern, carrier, carrier x pattern, and trigger cell the", None),
        ("audit is built on, plus the cross-corpus comparisons and within-corpus", None),
        ("inconsistency summary. Source = committed, gate-frozen ledgers only; no", None),
        ("recomputation. Regenerate with: python -m src reproduce statistics --output DIR.", None),
        ("", None),
        ("How to read positive_rate: per-cell share of RELEASED POSITIVES among firings.", Font(bold=True)),
        ("It is NOT precision -- the non-positives are not exhaustive true negatives.", None),
        ("", None),
        ("Cells with support < 5 are the excluded long tail (not assessed in the paper);", None),
        ("they are kept here with assessable=no so the whole distribution is visible.", None),
        ("", None),
        ("The 'reading' column spells each pattern out over placeholders A and B", None),
        ("(e.g. \"A bind B\", \"A interact with B\", \"A is an activator of B\"). Literal", None),
        ("corpus tokens are upright; inferred scaffolding words (is a / of / and) are italic.", None),
        ("", None),
        ("Sheets:", Font(bold=True)),
        ("  Patterns        pattern cells: family|shape|trigger|case|voice + reading x corpus", None),
        ("  Carriers        carrier-head cells x corpus (carrier_suspect flags parse-noise heads)", None),
        ("  CarrierPattern  carrier x pattern interaction cells", None),
        ("  Triggers        trigger-word cells x corpus", None),
        ("  CrossCorpus     shared-construction corpus_1 vs corpus_2 Fisher + BH (significant=q<0.05)", None),
        ("  WithinCorpus    per-corpus mixed/isolated-contrary summary", None),
        ("  CorePhysCells   core-physical within-corpus cells, carrier-stratified", None),
        ("", None),
        ("Eyeball guide for bug-hunting:", Font(bold=True)),
        ("  - Carriers with carrier_suspect=YES (tokens like '-', '3', punctuation) are", None),
        ("    prime candidates for a parse/segmentation bug.", None),
        ("  - A support>=5 pattern cell at positive_rate 0.0 or 1.0 is worth a sanity look.", None),
        ("  - CrossCorpus rows highlighted orange are the significant divergences.", None),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        if font:
            cell.font = font

    # ---- Patterns ----
    ws = wb.create_sheet("Patterns")
    headers = ["corpus", "family", "shape", "trigger", "case", "voice", "reading",
               "support", "positives", "negatives", "positive_rate",
               "wilson_low", "wilson_high", "assessable"]
    rows = []
    for row in sorted(by_dim.get("pattern", []), key=lambda r: -int(r["support"])):
        fam, shape, trig, case, voice = _split_pattern(row["key"])
        support = int(row["support"])
        rows.append([row["corpus"], fam, shape, trig, case, voice,
                     _readable_pattern(fam, shape, trig, case, voice), support,
                     _num(row["gold_positive"]), _num(row["gold_negative"]),
                     round(float(row["cpr"]), 4), round(float(row["wilson_low"]), 4),
                     round(float(row["wilson_high"]), 4),
                     "yes" if support >= ASSESSABLE_SUPPORT else "no"])
    _write_grid(ws, headers, rows,
                f"Pattern cells ({len(rows)} rows; {sum(1 for r in rows if r[-1]=='yes')} assessable, support>={ASSESSABLE_SUPPORT})",
                widths=[10, 11, 9, 14, 8, 9, 24, 9, 9, 9, 12, 10, 10, 10])

    # ---- Carriers ----
    ws = wb.create_sheet("Carriers")
    headers = ["corpus", "carrier_head", "support", "positives", "negatives",
               "positive_rate", "wilson_low", "wilson_high", "assessable", "carrier_suspect"]
    rows = []
    for row in sorted(by_dim.get("carrier", []), key=lambda r: -int(r["support"])):
        support = int(row["support"])
        rows.append([row["corpus"], row["key"], support,
                     _num(row["gold_positive"]), _num(row["gold_negative"]),
                     round(float(row["cpr"]), 4), round(float(row["wilson_low"]), 4),
                     round(float(row["wilson_high"]), 4),
                     "yes" if support >= ASSESSABLE_SUPPORT else "no",
                     "YES" if _carrier_suspect(row["key"]) else ""])

    def _carrier_fill(rec):
        return SUSPECT_FILL if rec[-1] == "YES" else None
    _write_grid(ws, headers, rows,
                f"Carrier-head cells ({len(rows)} rows; suspect heads highlighted)",
                widths=[10, 22, 9, 9, 9, 12, 10, 10, 10, 13], row_fill=_carrier_fill)

    # ---- CarrierPattern ----
    ws = wb.create_sheet("CarrierPattern")
    headers = ["corpus", "carrier_head", "pattern_key", "support", "positives",
               "negatives", "positive_rate", "assessable"]
    rows = []
    for row in sorted(by_dim.get("carrier_pattern", []), key=lambda r: -int(r["support"])):
        support = int(row["support"])
        carrier, _, patt = row["key"].partition(" | ")
        rows.append([row["corpus"], carrier, patt, support,
                     _num(row["gold_positive"]), _num(row["gold_negative"]),
                     round(float(row["cpr"]), 4),
                     "yes" if support >= ASSESSABLE_SUPPORT else "no"])
    _write_grid(ws, headers, rows, f"Carrier x pattern cells ({len(rows)} rows)",
                widths=[10, 18, 34, 9, 9, 9, 12, 10])

    # ---- Triggers ----
    ws = wb.create_sheet("Triggers")
    headers = ["corpus", "trigger", "support", "positives", "negatives",
               "positive_rate", "wilson_low", "wilson_high", "assessable"]
    rows = []
    for row in sorted(by_dim.get("trigger", []), key=lambda r: -int(r["support"])):
        support = int(row["support"])
        rows.append([row["corpus"], row["key"], support,
                     _num(row["gold_positive"]), _num(row["gold_negative"]),
                     round(float(row["cpr"]), 4), round(float(row["wilson_low"]), 4),
                     round(float(row["wilson_high"]), 4),
                     "yes" if support >= ASSESSABLE_SUPPORT else "no"])
    _write_grid(ws, headers, rows, f"Trigger cells ({len(rows)} rows)",
                widths=[10, 18, 9, 9, 9, 12, 10, 10, 10])

    # ---- CrossCorpus ----
    ws = wb.create_sheet("CrossCorpus")
    cross = _read(CROSS)
    headers = ["dimension", "family", "shape", "trigger", "case", "voice",
               "corpus_1", "cpr_1", "support_1", "corpus_2", "cpr_2", "support_2",
               "cpr_delta_abs", "p_value", "q_value", "significant"]
    rows = []
    n_sig = 0
    for row in sorted(cross, key=lambda r: float(r["q_value"])):
        fam, shape, trig, case, voice = _split_pattern(row["key"])
        sig = float(row["q_value"]) < 0.05
        n_sig += int(sig)
        rows.append([row["dimension"], fam, shape, trig, case, voice,
                     row["corpus_1"], round(float(row["cpr_1"]), 4), int(row["support_1"]),
                     row["corpus_2"], round(float(row["cpr_2"]), 4), int(row["support_2"]),
                     round(float(row["cpr_delta_abs"]), 4),
                     float(row["p_value"]), float(row["q_value"]),
                     "YES" if sig else ""])

    def _sig_fill(rec):
        return SIG_FILL if rec[-1] == "YES" else None
    _write_grid(ws, headers, rows,
                f"Cross-corpus shared constructions ({len(rows)} comparisons; {n_sig} significant at q<0.05)",
                widths=[10, 10, 9, 13, 7, 8, 9, 8, 9, 9, 8, 9, 12, 11, 11, 11],
                row_fill=_sig_fill)

    # ---- WithinCorpus ----
    ws = wb.create_sheet("WithinCorpus")
    within = _read(WITHIN)
    headers = list(within[0].keys()) if within else []
    rows = [[_num(v) for v in r.values()] for r in within]
    _write_grid(ws, headers, rows, f"Within-corpus inconsistency summary ({len(rows)} rows)",
                widths=[12, 18, 14, 14, 14, 14, 14, 18, 12])

    # ---- CorePhysCells ----
    ws = wb.create_sheet("CorePhysCells")
    core = _read(CORE)
    headers = list(core[0].keys()) if core else []
    rows = [[_num(v) for v in r.values()] for r in core]
    _write_grid(ws, headers, rows, f"Core-physical within-corpus cells ({len(rows)} rows)",
                widths=[10, 34, 8, 9, 8, 10, 9, 9, 9, 9, 8, 18, 20, 12, 16, 16])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path,
                        default=EXPERIMENT_DIR / 'appendix/pattern-carrier-review.xlsx')
    args = parser.parse_args(argv)
    for src in (CPR, CROSS, WITHIN, CORE):
        if not src.exists():
            sys.stderr.write(f"missing ledger: {src}\n")
            return 1
    out = build(args.out)
    print(f"wrote {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
